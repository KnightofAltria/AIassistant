"""Generate the Word counterpart of README.md; run only when content changes."""
import re
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
ROOT=Path(__file__).resolve().parent
def font(run,name='微软雅黑',size=10):
    run.font.name=name;run.font.size=Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)
    run.font.color.rgb=RGBColor(0,0,0)
def build():
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21.59);sec.page_height=Cm(27.94)
    sec.left_margin=sec.right_margin=Cm(2.1);sec.top_margin=sec.bottom_margin=Cm(1.8)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='微软雅黑';style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑');style.font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    normal=doc.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(6)
    code=False
    for line in (ROOT/'README.md').read_text(encoding='utf-8').splitlines():
        if line.startswith('```'):code=not code;continue
        if not line.strip():continue
        plain=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1（\2）',line).replace('**','').replace('`','')
        if line.startswith('# '):p=doc.add_paragraph(style='Title');font(p.add_run(plain[2:]),size=21)
        elif line.startswith('## '):p=doc.add_paragraph(style='Heading 1');font(p.add_run(plain[3:]),size=14)
        elif code:p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);font(p.add_run(plain),'Consolas',9)
        elif line.startswith('- '):p=doc.add_paragraph(style='List Bullet');font(p.add_run(plain[2:]))
        else:p=doc.add_paragraph();font(p.add_run(plain))
    footer=sec.footer.paragraphs[0];footer.alignment=2;font(footer.add_run('CSLA PC 版  |  '),size=8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.save(ROOT/'README.docx')
if __name__=='__main__':build()
