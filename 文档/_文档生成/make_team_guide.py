#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《队友协作指南》，同一份内容同时输出 Markdown 与 Word。

    python make_team_guide.py --out-dir 文档

输出：
    文档/队友协作指南.md      （GitHub 网页可直接阅读）
    文档/队友协作指南.docx    （可以直接发微信/QQ 给同学）
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x3B, 0x63)
GREY = RGBColor(0x66, 0x66, 0x66)

REPO_SSH = "git@github.com:KnightofAltria/AIassistant.git"
REPO_WEB = "https://github.com/KnightofAltria/AIassistant"

# ---------------------------------------------------------------- 内容定义
BLOCKS: list[tuple] = [
    ("h1", "队友协作指南"),
    ("sub", "计算机课程学习助手（CSLA）· 北京理工大学大四卓越工程实践"),

    ("h2", "一、五分钟上手"),
    ("bullets", [
        "把 GitHub 用户名发给组长，收到「已邀请」后接受邀请",
        "生成自己的 SSH 密钥（一条命令），把公钥加到你自己的 GitHub 账号",
        f"克隆仓库：`git clone {REPO_SSH}`",
        "干活前先 `git pull`，开自己的分支改，改完 `git push` 并开 Pull Request",
        "每天收工前确认当天的修改日志已提交",
    ]),

    ("h2", "二、你会用到的信息"),
    ("table", (
        ["项目", "内容"],
        [
            ["仓库地址（SSH）", REPO_SSH],
            ["仓库地址（网页）", REPO_WEB],
            ["主分支", "main（受保护，不直接提交）"],
            ["你的权限", "Write：可以推送、可以开 Pull Request"],
            ["要装的软件", "Git 2.x；写文档推荐 VS Code 或 Typora"],
            ["不需要装", "Word（能看 .docx 即可）、Python（除非你改生成脚本）"],
        ],
        [3.6, 12.6],
    )),

    ("h2", "三、关于公钥：每人一把，不能共用"),
    ("note", "公钥相当于身份证，**每个人必须用自己的**。组长那把只属于组长那台电脑——"
             "你把它加到你的账号也没用，因为配对的私钥不在你机器上。"),
    ("p", "正确做法：自己生成一把，把公钥贴到自己的 GitHub 账号。"),
    ("code", """ssh-keygen -t ed25519 -C "你的邮箱"
# 一路回车即可（密码可以留空）
# 公钥文件位置：
#   Windows     C:\\Users\\你的用户名\\.ssh\\id_ed25519.pub
#   macOS/Linux ~/.ssh/id_ed25519.pub"""),
    ("p", "公钥内容长这样，**一整行**，以 `ssh-ed25519` 开头："),
    ("code", "ssh-ed25519 AAAAC3NzaC1lZDI1NTE5AAAAIXXXX...（后面还有一长串）你的邮箱"),
    ("p", "复制这个 `.pub` 文件的**全部内容**，然后按顺序做："),
    ("bullets", [
        "打开 https://github.com/settings/ssh/new",
        "Title 随便填，比如「我的笔记本」",
        "Key 粘贴刚才复制的那一行",
        "Key type 选 **Authentication Key**（不是 Signing Key）",
        "点 Add SSH key 保存",
        "在本机验证：`ssh -T git@github.com`",
        "看到 `Hi 你的用户名! You've successfully authenticated` 就成功了",
    ]),
    ("note", "私钥文件 `id_ed25519`（**没有 .pub 后缀**）绝对不要发给任何人，"
             "也不要贴到群里、不要让 AI 助手读它。"),

    ("h2", "四、为什么不用 HTTPS + 密码/令牌"),
    ("p", "校园网和不少家庭宽带直连 `github.com:443` 会被阻断，典型表现是："),
    ("bullets", [
        "`git clone` 卡 20 秒左右后报 `Failed to connect to github.com port 443`",
        "或者 `Recv failure: Connection was reset`",
        "这时即使配了 PAT 令牌，照样推不上去",
    ]),
    ("p", "而 SSH（端口 22）一般是通的，所以全组统一走 SSH，省掉一堆麻烦。"),
    ("p", "万一 SSH 22 也不通，可以用 443 端口上的 SSH：在 `C:\\Users\\你的用户名\\.ssh\\config` 里写"),
    ("code", """Host github.com
  HostName ssh.github.com
  Port 443
  User git"""),

    ("h2", "五、第一次克隆仓库"),
    ("code", """git config --global user.name  "你的名字"
git config --global user.email "你的邮箱"
git clone git@github.com:KnightofAltria/AIassistant.git
cd AIassistant"""),
    ("note", "那两条 config 一定要填你自己的，否则提交记录显示成别人的名字，"
             "课程评「个人贡献」时会直接出问题。"),

    ("h2", "六、日常工作流"),
    ("p", "**不要直接在 main 上改**。每次干活走这个循环："),
    ("code", """git switch main
git pull                              # 1. 先同步最新
git switch -c docs/需求规格            # 2. 开自己的分支（docs/ code/ fix/ 开头）
# ...改文件...
git status                            # 3. 看清改了哪些
git add -A
git commit -m "docs: 补充需求规格第 3 章"   # 4. 提交
git push -u origin docs/需求规格        # 5. 推送"""),
    ("p", "推完 GitHub 页面会出现 **Compare & pull request** 按钮：点进去 → 写清楚改了什么 → "
          "至少一位队友 Review → Merge。"),
    ("p", "合并后回本地清理："),
    ("code", """git switch main
git pull
git branch -d docs/需求规格"""),

    ("h2", "七、提交信息怎么写"),
    ("table", (
        ["类型", "用在哪", "例子"],
        [
            ["feat", "新功能", "feat: 实现课程知识检索接口"],
            ["fix", "修 bug", "fix: 修正测试用例编号重复"],
            ["docs", "文档", "docs: 补充需求规格第 3 章"],
            ["refactor", "重构，不改功能", "refactor: 拆分文档生成脚本"],
            ["chore", "杂务（配置、依赖）", "chore: 更新 .gitignore"],
            ["test", "测试", "test: 增加登录用例"],
        ],
        [2.2, 5.0, 9.0],
    )),

    ("h2", "八、冲突了怎么办"),
    ("p", "两个人改了同一个文件的同一段，`git pull` 会提示 CONFLICT，文件里会出现："),
    ("code", """<<<<<<< HEAD
你写的内容
=======
队友写的内容
>>>>>>> origin/main"""),
    ("bullets", [
        "打开文件，保留要留下的内容，把 `<<<<<<<`、`=======`、`>>>>>>>` 三行标记全部删掉",
        "`git add 冲突的文件`",
        "`git commit`（如果用的是 rebase，则 `git rebase --continue`）",
        "`git push`",
    ]),
    ("p", "减少冲突最有效的办法不是记命令，而是：**一个文件只由一个人负责**；"
          "每天开工先 pull；长文档按章节拆成多个文件。"),

    ("h2", "九、仓库结构（东西放哪）"),
    ("code", """AIassistant/    项目代码
文档/           所有文档资料（课程要求、交付文档初稿、设计方案、生成脚本）
日志/           每日修改日志，一天一个 Word 文件
AGENTS.md       协作约定（给 AI 助手读的常驻规则，别删）
README.md       项目说明"""),
    ("bullets", [
        "新文档放 `文档/` 下对应的子目录，不要丢在仓库根目录",
        "不要提交 `__pycache__/`、`~$*.docx`（Word 临时锁文件）、`.zip`",
        "空文件夹 Git 不记录，要保留就在里面放一个 `.gitkeep`",
    ]),

    ("h2", "十、每日修改日志（全班统一要求）"),
    ("bullets", [
        "每天在 `日志/` 下有一份 `YYYY-MM-DD 修改日志.docx`",
        "由脚本生成，数据取自 git，**不手写、不编造**",
        "一天一份；同一天继续干活就更新那一份，不要新建第二份",
        "**当天谁干活谁负责生成**，别人不要改同一份——`.docx` 是二进制，Git 无法合并",
    ]),
    ("code", """python 文档/_文档生成/make_daily_log.py --repo . --date 2026-10-25 \\
  --note "今天做了什么" --problem "遇到的坑和解决办法" --todo "下一步\""""),

    ("h2", "十一、这些千万不要做"),
    ("bullets", [
        "不要 `git push --force`，会覆盖队友的提交",
        "不要 `git reset --hard` 之后再推送",
        "不要把私钥、密码、令牌提交进仓库——推上去就等于公开",
        "不要把 `node_modules/`、`.venv/` 这类依赖目录提交",
        "不要几天不 pull，攒一大堆再推",
    ]),

    ("h2", "十二、命令速查"),
    ("table", (
        ["想干什么", "命令"],
        [
            ["看当前改了什么", "git status"],
            ["看具体改动内容", "git diff"],
            ["看提交历史", "git log --oneline --graph"],
            ["拉最新代码", "git pull"],
            ["提交", 'git add -A && git commit -m "..."'],
            ["推送到远端", "git push"],
            ["撤销某文件的修改", "git restore 文件名"],
            ["撤销还没推的提交（保留改动）", "git reset --soft HEAD~1"],
            ["撤销已推送的提交", "git revert 提交哈希"],
            ["看某文件的历史", "git log --follow 文件名"],
            ["看某一行是谁写的", "git blame 文件名"],
            ["查看所有分支", "git branch -a"],
            ["切换分支", "git switch 分支名"],
            ["新分支第一次推送", "git push -u origin 分支名"],
        ],
        [6.0, 10.2],
    )),

    ("h2", "十三、人员与分工（请自行填写）"),
    ("table", (
        ["姓名", "角色", "GitHub 用户名", "负责模块"],
        [
            ["", "组长", "KnightofAltria", "仓库管理、文档整合"],
            ["", "组员", "待填", "待填"],
            ["", "组员", "待填", "待填"],
        ],
        [3.2, 2.4, 4.6, 6.0],
    )),

    ("h2", "十四、上手自检清单"),
    ("checklist", [
        "已收到仓库邀请，并在 GitHub 上点击了接受",
        "已生成自己的 SSH 密钥",
        "已把公钥加到自己的 GitHub（类型为 Authentication Key）",
        "`ssh -T git@github.com` 返回 You've successfully authenticated",
        "已克隆仓库到本地",
        "已设置自己的 user.name 和 user.email",
        "成功推送过一次分支，并开出 Pull Request",
        "知道每天的修改日志由谁生成、怎么生成",
    ]),
]


# ---------------------------------------------------------------- Word 渲染
def add_run(paragraph, text, font=BODY_FONT, size=10.5, bold=False, color=None):
    run = paragraph.add_run(text)
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def add_rich(paragraph, text, font=BODY_FONT, size=10.5, color=None):
    """把 **粗体** 和 `代码` 渲染成对应格式。"""
    import re
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
    for piece in pattern.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            add_run(paragraph, piece[2:-2], font=font, size=size, bold=True, color=color)
        elif piece.startswith("`") and piece.endswith("`"):
            add_run(paragraph, piece[1:-1], font=MONO_FONT, size=size - 1,
                    color=RGBColor(0xB0, 0x30, 0x60))
        else:
            add_run(paragraph, piece, font=font, size=size, color=color)


def shade_para(paragraph, fill="F2F2F2"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def shade_cell(cell, fill="DCE6F1"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def render_docx(path: Path):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.4)
    s.top_margin = s.bottom_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), BODY_FONT)

    for block in BLOCKS:
        kind = block[0]

        if kind == "h1":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run(p, block[1], font=HEAD_FONT, size=22, bold=True, color=ACCENT)

        elif kind == "sub":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run(p, block[1], size=9.5, color=GREY)
            doc.add_paragraph()

        elif kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            add_run(p, block[1], font=HEAD_FONT, size=14, bold=True, color=ACCENT)

        elif kind == "p":
            add_rich(doc.add_paragraph(), block[1])

        elif kind == "bullets":
            for item in block[1]:
                p = doc.add_paragraph(style="List Bullet")
                add_rich(p, item)

        elif kind == "checklist":
            for item in block[1]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(4)
                add_run(p, "☐  ", size=12, color=ACCENT)
                add_rich(p, item)

        elif kind == "note":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            shade_para(p, "FFF4D6")
            add_run(p, "提示　", font=HEAD_FONT, size=10, bold=True,
                    color=RGBColor(0xB0, 0x6A, 0x00))
            add_rich(p, block[1], size=10)

        elif kind == "code":
            for line in block[1].split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                shade_para(p, "F2F2F2")
                add_run(p, line if line else " ", font=MONO_FONT, size=9,
                        color=RGBColor(0x22, 0x33, 0x44))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

        elif kind == "table":
            headers, rows, widths = block[1]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for i, w in enumerate(widths):
                table.columns[i].width = Cm(w)
            for i, text in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell.width = Cm(widths[i])
                cell.text = ""
                add_run(cell.paragraphs[0], text, font=HEAD_FONT, size=9.5, bold=True)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade_cell(cell)
            for row in rows:
                cells = table.add_row().cells
                for i, text in enumerate(row):
                    cells[i].width = Cm(widths[i])
                    cells[i].text = ""
                    add_rich(cells[i].paragraphs[0], text, size=9)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, f"仓库：{REPO_WEB}", size=8.5, color=GREY)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


# ---------------------------------------------------------------- Markdown 渲染
def render_markdown(path: Path):
    out: list[str] = []
    for block in BLOCKS:
        kind = block[0]
        if kind == "h1":
            out.append(f"# {block[1]}\n")
        elif kind == "sub":
            out.append(f"> {block[1]}\n")
        elif kind == "h2":
            out.append(f"\n## {block[1]}\n")
        elif kind == "p":
            out.append(f"{block[1]}\n")
        elif kind == "bullets":
            out.extend(f"- {item}" for item in block[1])
            out.append("")
        elif kind == "checklist":
            out.extend(f"- [ ] {item}" for item in block[1])
            out.append("")
        elif kind == "note":
            out.append(f"> **提示**　{block[1]}\n")
        elif kind == "code":
            out.append("```")
            out.append(block[1])
            out.append("```\n")
        elif kind == "table":
            headers, rows, _ = block[1]
            out.append("| " + " | ".join(headers) + " |")
            out.append("| " + " | ".join("---" for _ in headers) + " |")
            for row in rows:
                out.append("| " + " | ".join(cell.replace("|", "\\|") for cell in row) + " |")
            out.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="生成队友协作指南（md + docx）")
    parser.add_argument("--out-dir", default="文档", help="输出目录")
    args = parser.parse_args()

    out_dir = Path(args.out_dir).resolve()
    md = out_dir / "队友协作指南.md"
    docx = out_dir / "队友协作指南.docx"
    render_markdown(md)
    render_docx(docx)
    print(f"已生成：{md}")
    print(f"已生成：{docx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
