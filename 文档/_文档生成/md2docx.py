#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把一份 Markdown 转成中文排版的 Word 文档。

支持：标题（#/##/###）、段落、无序列表、有序列表、表格、引用块、代码块、分割线、
行内 **加粗** 与 `等宽`。用于把交付文档的 .md 版本生成对应的 .docx。

    python md2docx.py 输入.md 输出.docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    """中文字体必须同时设 ascii 与 eastAsia，否则可能渲染成方框。"""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text: str, size: float = 10.5, base_bold: bool = False) -> None:
    """处理 **加粗** 与 `等宽` 两种行内标记。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, BODY_FONT, size, True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, MONO_FONT, size - 0.5, base_bold)
        else:
            run = paragraph.add_run(part)
            set_font(run, BODY_FONT, size, base_bold)


def setup(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(4)


def shade(cell, hex_color: str = "EDF1F8") -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear",
                                         qn("w:color"): "auto",
                                         qn("w:fill"): hex_color})
    tcpr.append(shd)


CONTENT_WIDTH_CM = 21.0 - 2.6 - 2.4     # 版心宽度，必须与 setup() 的页边距一致


def _set_table_width(table, total_cm: float, weights: list[float]) -> None:
    """显式写死表格与各列宽度，避免 Word/LibreOffice 自动扩展导致超出页面。"""
    from docx.oxml import OxmlElement
    tblpr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(total_cm / 2.54 * 1440)))
    tblpr.append(tblw)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)

    total = sum(weights) or 1.0
    for idx, col in enumerate(table.columns):
        width = Cm(total_cm * weights[idx] / total)
        col.width = width
        for cell in col.cells:
            cell.width = width


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 按每列最长内容估权重，限制在 1~3 之间，避免某列被挤成一条缝
    weights = []
    for j in range(cols):
        longest = max(len(r[j]) if j < len(r) else 0 for r in rows)
        weights.append(min(3.0, max(1.0, longest / 6 + 0.6)))
    _set_table_width(table, CONTENT_WIDTH_CM, weights)

    # 跨页时重复表头
    trpr = table.rows[0]._tr.get_or_add_trPr()
    header = trpr.makeelement(qn("w:tblHeader"), {qn("w:val"): "true"})
    trpr.append(header)
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            text = row[j] if j < len(row) else ""
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1.15
            add_inline(para, text, size=9.5, base_bold=(i == 0))
            if i == 0:
                shade(cell)
    document.add_paragraph()


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    setup(document)

    i = 0
    first_heading_done = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for code_line in block:
                para = document.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.1
                run = para.add_run(code_line)
                set_font(run, MONO_FONT, 9)
            document.add_paragraph()
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            rows = []
            header = [c.strip() for c in stripped.strip("|").split("|")]
            rows.append(header)
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(document, rows)
            continue

        # 分割线
        if stripped in ("---", "***", "___"):
            para = document.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            ppr = para._p.get_or_add_pPr()
            border = ppr.makeelement(qn("w:pBdr"), {})
            bottom = border.makeelement(qn("w:bottom"), {
                qn("w:val"): "single", qn("w:sz"): "6",
                qn("w:space"): "1", qn("w:color"): "CCCCCC"})
            border.append(bottom)
            ppr.append(border)
            i += 1
            continue

        # 标题
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and not first_heading_done:
                para = document.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(12)
                para.paragraph_format.space_after = Pt(18)
                run = para.add_run(text)
                set_font(run, HEAD_FONT, 19, True)
                first_heading_done = True
            else:
                para = document.add_paragraph()
                para.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
                para.paragraph_format.space_after = Pt(6)
                para.paragraph_format.keep_with_next = True
                run = para.add_run(text)
                set_font(run, HEAD_FONT, {1: 14, 2: 12.5}.get(level, 11.5), True)
                run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            para = document.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            para.paragraph_format.space_after = Pt(6)
            run = para.add_run(text)
            set_font(run, BODY_FONT, 10)
            run.font.color.rgb = RGBColor(0x55, 0x5F, 0x6D)
            i += 1
            continue

        # 列表
        if re.match(r"^[-*]\s+", stripped):
            para = document.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue
        if re.match(r"^\d+[.)]\s+", stripped):
            para = document.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, re.sub(r"^\d+[.)]\s+", "", stripped))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        para = document.add_paragraph()
        para.paragraph_format.first_line_indent = Cm(0.0)
        add_inline(para, stripped)
        i += 1

    document.save(docx_path)


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    md_path, docx_path = Path(sys.argv[1]), Path(sys.argv[2])
    convert(md_path, docx_path)
    print(f"已生成：{docx_path}  （{docx_path.stat().st_size / 1024:.1f} KB）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
