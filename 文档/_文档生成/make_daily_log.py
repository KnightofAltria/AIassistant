#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成「每日修改日志」Word 文档。

数据全部来自 git，不手写、不编造：
    python make_daily_log.py --repo <仓库根> --date 2026-10-24 \\
        --note "导入全部资料" --problem "HTTPS 被阻断" --todo "邀请队友"

输出默认写到 <仓库根>/日志/<YYYY-MM-DD> 修改日志.docx
"""

from __future__ import annotations

import argparse
import datetime as dt
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"
SEP = "\x1f"                      # 分隔符，避免提交信息里的空格干扰解析
TYPES = ("feat", "fix", "docs", "chore", "refactor", "test", "style", "perf")


# ---------------------------------------------------------------- git 数据层
def git(repo: Path, *args: str) -> str:
    proc = subprocess.run(
        ["git", "-C", str(repo), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if proc.returncode != 0:
        raise RuntimeError(f"git {' '.join(args)} 执行失败：{proc.stderr.strip()}")
    return proc.stdout


def collect_commits(repo: Path, day: str) -> list[dict]:
    """取当天 00:00 至 24:00 的提交（含文件变动统计）。"""
    raw = git(
        repo, "log",
        f"--since={day} 00:00:00", f"--until={day} 23:59:59",
        "--reverse", f"--pretty=format:%h{SEP}%ad{SEP}%an{SEP}%s",
        "--date=format:%H:%M",
    )
    commits = []
    for line in raw.splitlines():
        if not line.strip():
            continue
        sha, when, author, subject = line.split(SEP, 3)
        stat = git(repo, "show", "--shortstat", "--format=", sha).strip()
        commits.append(
            {"sha": sha, "time": when, "author": author, "subject": subject,
             "stat": summarize_stat(stat)}
        )
    return commits


def summarize_stat(stat: str) -> str:
    """把 '99 files changed, 20257 insertions(+)' 压成 '99 文件 / +20257'。"""
    if not stat:
        return "—"
    files = ins = dele = 0
    for part in stat.split(","):
        part = part.strip()
        num = "".join(ch for ch in part.split(" ")[0] if ch.isdigit())
        if "file" in part:
            files = int(num or 0)
        elif "insertion" in part:
            ins = int(num or 0)
        elif "deletion" in part:
            dele = int(num or 0)
    bits = [f"{files} 文件"]
    if ins:
        bits.append(f"+{ins}")
    if dele:
        bits.append(f"-{dele}")
    return " / ".join(bits)


def commit_type(subject: str) -> str:
    if ":" in subject:
        head = subject.split(":", 1)[0].strip().lower()
        if head in TYPES:
            return head
    return "其他"


def subject_body(subject: str) -> str:
    """去掉 'chore: ' 这类前缀，类型已经在单独的列里。"""
    if ":" in subject:
        head, rest = subject.split(":", 1)
        if head.strip().lower() in TYPES:
            return rest.strip()
    return subject


def repo_status(repo: Path) -> dict:
    total = int(git(repo, "rev-list", "--count", "HEAD").strip())
    tracked = len([l for l in git(repo, "ls-tree", "-r", "--name-only", "HEAD").splitlines() if l])
    branch = git(repo, "rev-parse", "--abbrev-ref", "HEAD").strip()
    head = git(repo, "rev-parse", "HEAD").strip()
    dirty = bool(git(repo, "status", "--porcelain").strip())
    remote_head = ""
    try:
        remote_head = git(repo, "rev-parse", f"origin/{branch}").strip()
    except RuntimeError:
        pass
    if not remote_head:
        sync = "未找到远端分支"
    elif dirty:
        sync = "本地有未提交改动"
    elif head == remote_head:
        sync = "与远端一致"
    else:
        sync = "与远端不一致（需推送或拉取）"
    size = sum(p.stat().st_size for p in repo.rglob("*") if p.is_file() and ".git" not in p.parts)
    return {"total": total, "tracked": tracked, "branch": branch, "head": head[:7],
            "dirty": dirty, "sync": sync, "size_mb": round(size / 1024 / 1024, 2)}


# ---------------------------------------------------------------- 排版工具
def set_style_font(style, name: str, size: float | None = None,
                   bold: bool | None = None, color: RGBColor | None = None) -> None:
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def add_run(paragraph, text: str, font: str = BODY_FONT, size: float = 10.5,
            bold: bool = False, color: RGBColor | None = None):
    run = paragraph.add_run(text)
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def shade(cell, fill: str = "DCE6F1") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def make_table(doc, headers: list[str], rows: list[list[str]], widths_cm: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # 固定布局：LibreOffice / Word 都以 tblGrid 的 gridCol 为准，
    # 只设 cell.width 会出现"每列等宽"的错版，必须同时写列宽。
    for i, width in enumerate(widths_cm):
        table.columns[i].width = Cm(width)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Cm(widths_cm[i])
        cell.text = ""
        add_run(cell.paragraphs[0], text, font=HEAD_FONT, size=9.5, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(cell)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = Cm(widths_cm[i])
            cells[i].text = ""
            add_run(cells[i].paragraphs[0], text, size=9.5)
    return table


def heading(doc, text: str, level: int = 1) -> None:
    para = doc.add_heading(level=level)
    add_run(para, text, font=HEAD_FONT, size=14 if level == 1 else 12, bold=True,
            color=RGBColor(0x1F, 0x3B, 0x63))


def bullet(doc, text: str) -> None:
    para = doc.add_paragraph(style="List Bullet")
    add_run(para, text, size=10.5)


# ---------------------------------------------------------------- 主流程
def build(repo: Path, day: str, out: Path, notes: list[str], problems: list[str],
          todos: list[str], author: str) -> Path:
    commits = collect_commits(repo, day)
    status = repo_status(repo)
    weekday = "一二三四五六日"[dt.date.fromisoformat(day).weekday()]

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.4)
    section.top_margin = section.bottom_margin = Cm(2.2)

    set_style_font(doc.styles["Normal"], BODY_FONT, 10.5)
    for name in ("Heading 1", "Heading 2", "Title"):
        if name in [s.name for s in doc.styles]:
            set_style_font(doc.styles[name], HEAD_FONT)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.4
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)

    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, f"{day} 修改日志", font=HEAD_FONT, size=20, bold=True,
            color=RGBColor(0x1F, 0x3B, 0x63))
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(sub, f"计算机课程学习助手（CSLA）　星期{weekday}　记录人：{author}",
            size=9.5, color=RGBColor(0x60, 0x60, 0x60))
    doc.add_paragraph()

    # 一、今日概述
    heading(doc, "一、今日概述")
    para = doc.add_paragraph()
    add_run(para, f"本日共产生 {len(commits)} 次提交，涉及 "
                  f"{sum(1 for c in commits)} 条记录，当前仓库已跟踪 "
                  f"{status['tracked']} 个文件、累计 {status['total']} 次提交。", size=10.5)
    for note in notes:
        bullet(doc, note)
    if not commits and not notes:
        bullet(doc, "本日无提交，也未改动仓库内容。")

    # 二、变更清单
    heading(doc, "二、变更清单")
    if commits:
        rows = [[c["time"], c["sha"], commit_type(c["subject"]),
                 subject_body(c["subject"]), c["stat"]]
                for c in commits]
        make_table(doc, ["时间", "提交", "类型", "说明", "文件变化"], rows,
                   [1.6, 1.9, 1.5, 7.5, 3.7])
        tip = doc.add_paragraph()
        add_run(tip, "数据来源：git log --since / --until + git show --shortstat。"
                     "生成本文档的那次提交不计入上表。",
                size=8.5, color=RGBColor(0x80, 0x80, 0x80))
    else:
        para = doc.add_paragraph()
        add_run(para, "本日无提交。", size=10.5)

    # 三、仓库当前状态
    heading(doc, "三、仓库当前状态")
    rows = [
        ["当前分支", status["branch"]],
        ["最新提交", status["head"]],
        ["提交总数", str(status["total"])],
        ["今日提交", str(len(commits))],
        ["跟踪文件数", str(status["tracked"])],
        ["工作区", "有未提交改动" if status["dirty"] else "干净"],
        ["与远端", status["sync"]],
        ["仓库体积", f"{status['size_mb']} MB（不含 .git）"],
    ]
    make_table(doc, ["项目", "状态"], rows, [4.0, 12.2])

    # 四、问题与处理
    heading(doc, "四、问题与处理")
    if problems:
        for item in problems:
            bullet(doc, item)
    else:
        para = doc.add_paragraph()
        add_run(para, "本日无阻塞性问题。", size=10.5)

    # 五、遗留与下一步
    heading(doc, "五、遗留与下一步")
    if todos:
        for item in todos:
            bullet(doc, item)
    else:
        para = doc.add_paragraph()
        add_run(para, "暂无记录。", size=10.5)

    # 页脚
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(footer, f"生成时间：{dt.datetime.now():%Y-%m-%d %H:%M}", size=8.5,
            color=RGBColor(0x88, 0x88, 0x88))

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="生成每日修改日志 Word 文档")
    parser.add_argument("--repo", default=".", help="仓库根目录")
    parser.add_argument("--date", default=dt.date.today().isoformat(), help="日志日期 YYYY-MM-DD")
    parser.add_argument("--out", default=None, help="输出文件路径")
    parser.add_argument("--note", action="append", default=[], help="今日概述要点，可多次")
    parser.add_argument("--problem", action="append", default=[], help="问题与处理，可多次")
    parser.add_argument("--todo", action="append", default=[], help="遗留与下一步，可多次")
    parser.add_argument("--author", default=None, help="记录人，默认取 git config user.name")
    args = parser.parse_args(argv)

    repo = Path(args.repo).resolve()
    if not (repo / ".git").exists():
        print(f"不是 git 仓库：{repo}", file=sys.stderr)
        return 2
    author = args.author or git(repo, "config", "user.name").strip() or "—"
    out = Path(args.out) if args.out else repo / "日志" / f"{args.date} 修改日志.docx"

    path = build(repo, args.date, out, args.note, args.problem, args.todo, author)
    print(f"已生成：{path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
